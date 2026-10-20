"""`openpyxl_csv` 阶段：CSV -> xlsx -> 读回 -> CSV（openpyxl 3.x）。

它是 `excel_com_csv` 的**可跑 fallback**，不是它的替代品（D1 / K3）：
- openpyxl 写数字用 `%.16g`（F5），Excel COM 走 Excel 自己的数字模型 → **损失形态不同**；
- 因此本阶段产出的差异**不得**归因到 `excel_15_digit_truncation`（M3 的
  `test_candidate_stage_not_in_run_degrades_to_unclassified` 就钉这条）。

写入侧的类型决策是我们自己的代码（`stages/common.cell_kind`）：整数形且无前导零才算数字，
`02114` / `00123-4567` / `1900-02-29` 一律按文本原样交给 openpyxl。这条启发式写死并测试钉住，
否则"到底是谁把 `02114` 变成 `2114`"就分辨不出是本阶段还是被测引擎。
"""

from __future__ import annotations

import csv
import io
import shutil
import tempfile
from pathlib import Path

from tripwire.pipeline.errors import StageFailed, StageUnavailable
from tripwire.pipeline.model import ColumnMeta, RuntimeMeta, StageResult
from tripwire.pipeline.stages.common import coerce, render_cell, rows_to_text

STAGE_NAME = "openpyxl_csv"


def _require_openpyxl():
    try:
        import openpyxl  # noqa: PLC0415 —— 可选依赖，必须函数内导入
    except ImportError as error:  # pragma: no cover - 依赖缺失路径
        raise StageUnavailable(STAGE_NAME, "openpyxl 未安装（pip install \"tripwire\"）") from error
    return openpyxl


def run(text: str) -> StageResult:
    """走一遍 xlsx 往返，返回读回后的**原始文本**（不回传 workbook，K1）。"""
    openpyxl = _require_openpyxl()

    records = [row for row in csv.reader(io.StringIO(text, newline=""))]
    if not records:
        return StageResult(text="", columns=(), runtime=RuntimeMeta(openpyxl_version=str(openpyxl.__version__)))

    header = records[0]
    workdir = Path(tempfile.mkdtemp(prefix="tripwire-openpyxl-"))
    xlsx = workdir / "book.xlsx"
    try:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row_index, record in enumerate(records, start=1):
            for col_index, cell_text in enumerate(record, start=1):
                sheet.cell(row=row_index, column=col_index).value = coerce(cell_text)
        try:
            workbook.save(xlsx)
        except Exception as error:
            raise StageFailed(STAGE_NAME, f"openpyxl 写 xlsx 失败: {error}") from error
        finally:
            workbook.close()

        try:
            loaded = openpyxl.load_workbook(xlsx, data_only=True)
        except Exception as error:
            raise StageFailed(STAGE_NAME, f"openpyxl 读回 xlsx 失败: {error}") from error
        try:
            sheet = loaded.active
            # max_row/max_col 会把"尾部全空"的行裁掉，这是 xlsx 的既有行为，如实反映
            rows = [
                tuple(render_cell(value) for value in row)
                for row in sheet.iter_rows(
                    min_row=1, max_row=sheet.max_row, max_col=sheet.max_column, values_only=True
                )
            ]
        finally:
            loaded.close()

        if not rows:
            raise StageFailed(STAGE_NAME, "读回的 xlsx 没有任何行")
        out_text = rows_to_text(rows)
        columns = tuple(ColumnMeta(name=str(name), declared_type=None) for name in rows[0])
        return StageResult(
            text=out_text,
            columns=columns,
            runtime=RuntimeMeta(openpyxl_version=str(openpyxl.__version__)),
        )
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


def available() -> tuple[bool, str]:
    try:
        _require_openpyxl()
    except StageUnavailable as error:
        return False, error.reason
    return True, ""


__all__ = ["STAGE_NAME", "available", "run"]
